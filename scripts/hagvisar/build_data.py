#!/usr/bin/env python3
"""Convert the Hagvísar xlsx workbooks from sedlabanki.is into the JSON the web app reads.

Usage:
    python scripts/hagvisar/build_data.py                 # download all workbooks and rebuild public/data
    python scripts/hagvisar/build_data.py --local DIR     # parse previously downloaded workbooks in DIR
    python scripts/hagvisar/build_data.py --check         # only ask the server whether anything changed

Outputs (default --out public/data):
    catalogue.json        figure metadata for every figure (loaded once by the app)
    figures/<id>.json     x values and series values for one figure (loaded on demand)
    manifest.json         per-workbook URL, ETag, Last-Modified and sha256 (used by --check)

--check exits with 0 when every workbook matches manifest.json and 10 when at least one
has changed (or is new), so an update service can run `--check || build`.
"""

from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import json
import math
import re
import statistics
import sys
import tempfile
import urllib.request
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent.parent
CHAPTERS = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X"]
USER_AGENT = "Mozilla/5.0 (compatible; betrihagvisar-data/1.0)"
EXIT_CHANGED = 10

META_LABELS = {"fs": "title", "ufs": "subtitle", "nm": "note", "h": "sources", "vá": "unit", "há": "unitRight", "ath": "comment", "aths": "comment"}
META_ROWS = {  # fallback when column A has no labels: 1-based row -> key; values sit in column B
    4: "title",
    5: "subtitle",
    6: "note",
    7: "sources",
    8: "unit",
    9: "unitRight",
    10: "comment",
}

warnings: list[str] = []


def warn(fig_id: str, msg: str) -> None:
    warnings.append(f"{fig_id}: {msg}")


# --------------------------------------------------------------------------- #
#  Download / check
# --------------------------------------------------------------------------- #


def http(url: str, method: str = "GET"):
    req = urllib.request.Request(url, method=method, headers={"User-Agent": USER_AGENT})
    return urllib.request.urlopen(req, timeout=60)


def download(url: str, dest: Path) -> dict:
    with http(url) as r:
        body = r.read()
        info = {"etag": r.headers.get("ETag"), "lastModified": r.headers.get("Last-Modified")}
    dest.write_bytes(body)
    info["sha256"] = hashlib.sha256(body).hexdigest()
    info["bytes"] = len(body)
    return info


def check(sources: dict, manifest_path: Path) -> int:
    old = json.loads(manifest_path.read_text("utf-8")) if manifest_path.exists() else {"workbooks": {}}
    changed = []
    for ch, url in sources["workbooks"].items():
        prev = old["workbooks"].get(ch)
        if not prev or prev.get("url") != url:
            changed.append(f"{ch}: new URL")
            continue
        with http(url, "HEAD") as r:
            etag, lm = r.headers.get("ETag"), r.headers.get("Last-Modified")
        if (etag and etag != prev.get("etag")) or (lm and lm != prev.get("lastModified")):
            changed.append(f"{ch}: ETag/Last-Modified changed")
    for line in changed:
        print(line)
    print("Changed." if changed else "Up to date.")
    return EXIT_CHANGED if changed else 0


# --------------------------------------------------------------------------- #
#  Cell helpers
# --------------------------------------------------------------------------- #


def clean_text(v) -> str | None:
    if v is None:
        return None
    s = re.sub(r"\s+", " ", str(v)).strip()
    return s or None


def to_number(v) -> float | None:
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return None if (isinstance(v, float) and (math.isnan(v) or math.isinf(v))) else float(v)
    s = str(v).strip().replace("−", "-").replace(",", ".")
    if re.fullmatch(r"-?\d+(\.\d+)?([eE]-?\d+)?", s):
        return float(s)
    return None


def round_value(v: float | None) -> float | int | None:
    if v is None:
        return None
    r = round(v, 6)
    return int(r) if r == int(r) and abs(r) < 1e15 else r


def blank(row) -> bool:
    return all(v is None or (isinstance(v, str) and not v.strip()) for v in row)


def label(v) -> str:
    if isinstance(v, dt.datetime):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    s = clean_text(v) or ""
    return f"'{s}" if re.fullmatch(r"\d\d", s) else s  # two-digit years as in the report: '91


# --------------------------------------------------------------------------- #
#  X-axis parsing
# --------------------------------------------------------------------------- #

MONTHS = {
    "jan": 1, "feb": 2, "mar": 3, "apr": 4, "maí": 5, "mai": 5, "jún": 6, "jun": 6, "júl": 7, "jul": 7,
    "ágú": 8, "ágú": 8, "agu": 8, "ág": 8, "sep": 9, "okt": 10, "nóv": 11, "nov": 11, "des": 12, "dec": 12,
}
FREQ_PER_YEAR = {1: "annual", 2: "semiannual", 4: "quarterly", 6: "bimonthly", 12: "monthly"}


def month_from(word: str) -> int | None:
    w = word.lower().strip(". ")
    for k in sorted(MONTHS, key=len, reverse=True):
        if w.startswith(k):
            return MONTHS[k]
    return None


def year_from(s: str) -> int | None:
    s = s.strip().strip("'\"’ ").strip()
    if re.fullmatch(r"(19|20)\d\d", s):
        return int(s)
    if re.fullmatch(r"\d\d", s):
        y = int(s)
        return 1900 + y if y >= 50 else 2000 + y
    return None


def parse_token(v):
    """Return (year, sub, per_year) where sub is 0-based within the year, or None.

    per_year is 4 for quarters, 12 for months, 6 for two-month periods; None when only the
    year is known. Returns ("date", datetime) for exact dates.
    """
    if v is None:
        return None
    if isinstance(v, dt.datetime):
        return ("date", v)
    if isinstance(v, (int, float)):
        if isinstance(v, int) or float(v).is_integer():
            n = int(v)
            if 19000101 <= n <= 21001231:  # yyyymmdd
                try:
                    return ("date", dt.datetime.strptime(str(n), "%Y%m%d"))
                except ValueError:
                    return None
            if 1900 <= n <= 2100:
                return (n, None, None)
            return None
        if 1900 <= v <= 2100:  # Excel artefacts such as 2025.1666…
            return (int(v), None, None)
        return None
    s = str(v).strip().strip("'\"’ ").strip()
    if not s:
        return None
    m = re.fullmatch(r"(\d)\.?\s*ársfj\.?\s*(\d{4})", s, re.I)
    if m:
        return (int(m.group(2)), int(m.group(1)) - 1, 4)
    m = re.fullmatch(r"(\d{4})\s*[:\-]\s*([1-4])", s)
    if m:
        return (int(m.group(1)), int(m.group(2)) - 1, 4)
    m = re.fullmatch(r"(\d{4})\s*(?:F|Q)\s*([1-4])", s, re.I)
    if m:
        return (int(m.group(1)), int(m.group(2)) - 1, 4)
    m = re.fullmatch(r"(?:Q|F)?([1-4])\s*(?:F|Q)?\s+(\d{4})", s, re.I)
    if m and re.search(r"[FQ]", s, re.I):
        return (int(m.group(2)), int(m.group(1)) - 1, 4)
    m = re.fullmatch(r"(\d{4})\s*M\s*(\d{1,2})", s, re.I)
    if m:
        return (int(m.group(1)), int(m.group(2)) - 1, 12)
    m = re.fullmatch(r"([A-Za-zÁáÍíÓóÚúÝýÞþÆæÖöÐðÉé]+)[.,]?\s*-\s*([A-Za-zÁáÍíÓóÚúÝýÞþÆæÖöÐðÉé]+)\.?\s+(\d{4})", s)
    if m:
        a = month_from(m.group(1))
        if a:
            return (int(m.group(3)), (a - 1) // 2, 6)
    m = re.fullmatch(r"(\d{1,2})\.(\d{1,2})\.(\d{4})", s)
    if m:
        d, mo, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
        return ("date", dt.datetime(y, mo, min(d, 28)))
    y = year_from(s)
    if y:
        return (y, None, None)
    return None


def period_start(year: int, sub: int, per_year: int) -> dt.date:
    month = 1 + sub * (12 // per_year)
    return dt.date(year, month, 1)


def infer_freq_from_dates(dates: list[dt.date]) -> str:
    gaps = [(b - a).days for a, b in zip(dates, dates[1:]) if b > a]
    if not gaps:
        return "annual"
    g = statistics.median(gaps)
    if g <= 4:
        return "daily"
    if g <= 10:
        return "weekly"
    if g <= 45:
        return "monthly"
    if g <= 75:
        return "bimonthly"
    if g <= 120:
        return "quarterly"
    if g <= 250:
        return "semiannual"
    return "annual"


def normalise_dates(dates: list[dt.date], freq: str) -> list[dt.date]:
    """Move monthly/quarterly/annual observations to the start of their period."""
    if freq in ("daily", "weekly"):
        return dates
    per = {"monthly": 12, "bimonthly": 6, "quarterly": 4, "semiannual": 12, "annual": 1}[freq]
    step = 12 // per
    out = [dt.date(d.year, ((d.month - 1) // step) * step + 1, 1) for d in dates]
    if len(set(out)) != len(out):
        return dates
    return out


def regular_sequence(tokens: list):
    """Fit one regular frequency to the tokens, anchored on explicit sub-year labels.

    Year cells next to quarter labels are sometimes off by one in the workbooks, so up to
    10% of year-only tokens may disagree. Returns (dates, frequency) or None.
    """
    explicit = [(i, t) for i, t in enumerate(tokens) if t and t[1] is not None]
    if explicit:
        pers = {t[2] for _, t in explicit}
        if len(pers) != 1:
            return None
        per = pers.pop()
        k0 = Counter(t[0] * per + t[1] - i for i, t in explicit).most_common(1)[0][0]
        bad = [i for i, t in explicit if t[0] * per + t[1] != k0 + i]
        if len(bad) > max(1, len(explicit) // 20):
            return None
    else:
        runs = []
        for t in tokens:
            if t is None:
                return None
            if runs and runs[-1][0] == t[0]:
                runs[-1][1] += 1
            else:
                runs.append([t[0], 1])
        inner = [n for _, n in runs[1:-1]] or [runs[0][1]]
        per = Counter(inner).most_common(1)[0][0]
        if per not in (1, 2, 4, 6, 12) or len(runs) < 2:
            return None
        idx = next((i for i, (_, n) in enumerate(runs) if n == per), None)
        if idx is None:
            return None
        before = sum(n for _, n in runs[:idx])
        k0 = runs[idx][0] * per - before
    years = [t[0] for t in tokens if t is not None and t[1] is None]
    fitted = [(k0 + i) // per for i, t in enumerate(tokens) if t is not None and t[1] is None]
    misses = sum(1 for y, f in zip(years, fitted) if y != f)
    if years and misses > max(1, len(years) // 10):
        return None
    dates = [period_start((k0 + i) // per, (k0 + i) % per, per) for i in range(len(tokens))]
    return dates, FREQ_PER_YEAR[per]


def mixed_sequence(tokens: list):
    """Annual history followed by sub-annual data (e.g. IV-1: yearly until 2003, then quarterly)."""
    if any(t is None or t[1] is not None for t in tokens):
        return None
    runs: list[list] = []
    for t in tokens:
        if runs and runs[-1][0] == t[0]:
            runs[-1][1] += 1
        else:
            runs.append([t[0], 1])
    years = [y for y, _ in runs]
    if len(set(years)) != len(years) or years != sorted(years):
        return None
    out = []
    for i, (y, n) in enumerate(runs):
        if n not in (1, 2, 4, 6, 12) and i != len(runs) - 1:
            return None
        per = n if n in (1, 2, 4, 6, 12) else runs[i - 1][1]
        if i == len(runs) - 1 and i > 0 and n < runs[i - 1][1]:
            per = runs[i - 1][1]
        out.extend((y, s, per) for s in range(n))
    pers = [p for _, _, p in out]
    # A long annual history followed by a few sub-annual points (year-to-date, latest
    # quarter) is drawn as categories in the report; keep it that way.
    if pers[0] == 1 and sum(1 for p in pers if p != 1) <= 3:
        return None
    dates = [period_start(y, s, p) for y, s, p in out]
    if any(b <= a for a, b in zip(dates, dates[1:])):
        return None
    tail = [p for p in pers if p != 1]
    per = Counter(tail or pers).most_common(1)[0][0]
    return dates, FREQ_PER_YEAR[per]


LETTER_MONTHS = list("JFMAMJJÁSOND")


def letters_monthly(a_col: list, b_col: list):
    """Year marker in column A on the January row, month initials in column B (VIII-35/36)."""
    if not all(isinstance(b, str) and len(b.strip()) == 1 for b in b_col):
        return None
    for i, (a, b) in enumerate(zip(a_col, b_col)):
        tok = parse_token(a)
        if tok and tok[0] != "date" and tok[1] is None and b.strip().upper() == "J":
            y0 = tok[0]
            dates = []
            for k in range(len(b_col)):
                m = k - i
                y, mo = y0 + (m // 12), m % 12
                dates.append(dt.date(y, mo + 1, 1))
                if b_col[k].strip().upper() != LETTER_MONTHS[mo]:
                    return None
            return dates, "monthly"
    return None


def resolve_x(a_col: list, b_col: list | None):
    """Return (xType, values, frequency). values are ISO dates for time, labels for category."""
    # 1. exact dates in column A (or in an unlabeled column B)
    for col in (a_col, b_col):
        if col is None:
            continue
        toks = [parse_token(v) for v in col]
        if toks and all(t and t[0] == "date" for t in toks):
            dates = [t[1].date() for t in toks]
            if any(b < a for a, b in zip(dates, dates[1:])):
                break
            if len(set(dates)) != len(dates):
                return "time-dedupe", [d.isoformat() for d in dates], infer_freq_from_dates(dates)
            freq = infer_freq_from_dates(dates)
            dates = normalise_dates(dates, freq)
            if len(set(dates)) == len(dates):
                return "time", [d.isoformat() for d in dates], freq
    # 2. month initials with year markers
    if b_col is not None:
        r = letters_monthly(a_col, b_col)
        if r:
            return "time", [d.isoformat() for d in r[0]], r[1]
    # 3. year / quarter / month tokens
    toks = [parse_token(v) for v in a_col]
    if toks and not any(t and t[0] == "date" for t in toks) and sum(t is None for t in toks) <= max(1, len(toks) // 20):
        r = regular_sequence(toks) or mixed_sequence(toks)
        if r:
            return "time", [d.isoformat() for d in r[0]], r[1]
    return "category", [label(v) for v in a_col], None


# --------------------------------------------------------------------------- #
#  Sheet parsing
# --------------------------------------------------------------------------- #

AXIS_RE = re.compile(r"\s*\((andhverfur\s+)?(v|h)\.\s*ás\)\s*", re.I)


def split_axis(name: str) -> tuple[str, str]:
    """'Gengi (h. ás)' -> ('Gengi', 'right'); an inverted axis ('andhverfur h. ás') -> 'right-inverse'."""
    m = AXIS_RE.search(name)
    if not m:
        return name, "left"
    side = "right" if m.group(2).lower() == "h" else "left"
    return (AXIS_RE.sub(" ", name).strip(), f"{side}-inverse" if m.group(1) else side)


def parse_sources(s: str | None) -> list[str]:
    if not s:
        return []
    s = re.sub(r"^\s*Heimild(ir)?\s*:\s*", "", s, flags=re.I).strip().rstrip(".").strip()
    parts = re.split(r",\s*|\s+og\s+", s)
    return [p.strip().rstrip(".") for p in parts if p.strip()]


def is_string(v) -> bool:
    return isinstance(v, str) and v.strip() != "" and to_number(v) is None


def find_data_start(rows: list[tuple], first: int = 10) -> int | None:
    """First row (0-based, after the metadata block) that starts the data block.

    A data row has an x value in column A (or a date in column B) and no text from column C
    on. Leading rows may have no values yet, so the block only has to contain a number
    somewhere further down.
    """
    # numbers_after[i]: is there a number in column B onward in row i or later?
    numbers_after = [False] * (len(rows) + 1)
    for i in range(len(rows) - 1, -1, -1):
        numbers_after[i] = numbers_after[i + 1] or any(to_number(v) is not None for v in (rows[i][1:] if rows[i] else ()))
    for i in range(first, len(rows)):
        row = rows[i]
        if not numbers_after[i]:
            return None
        if not row or blank(row):
            continue
        x = row[0]
        if x is None and len(row) > 1 and isinstance(row[1], dt.datetime) and not any(isinstance(v, dt.datetime) for v in row[2:]):
            x = row[1]
        if x is None or any(is_string(v) for v in row[2:]):
            continue
        if is_year_header(row, rows[i + 1 :]):
            continue
        return i
    return None


def is_year_header(row, rest) -> bool:
    """A row of plain years above the data (III-19) is a header, not data."""
    nums = [v for v in row[1:] if v is not None]
    if not nums or not all(isinstance(v, int) and 1900 <= v <= 2100 for v in nums):
        return False
    nxt = next((r for r in rest if r and not blank(r)), None)
    return bool(nxt) and any(isinstance(v, float) and not v.is_integer() for v in nxt[1:])


def transpose_block(header: list, data: list[list]) -> tuple[list, list[list]]:
    """Series in rows and dates across (IX-2): turn into dates down, series across."""
    cols = [c for c in range(1, len(header)) if isinstance(header[c], dt.datetime)]
    new_header = [None, None] + [r[0] for r in data]
    new_data = [[header[c], None] + [r[c] for r in data] for c in cols]
    return new_header, new_data


UNIT_RE = re.compile(r"^(%|ma\.|mö\.|ma |m\.kr|þús|vísitala|fjöldi|breyting|hlutfall|usd|eur|kr\.|punktar|prósent|ár$|mánuðir|dagar|% af|staða|vaxtaálag$)", re.I)


def looks_like_unit(text: str) -> bool:
    return len(text) <= 60 and bool(UNIT_RE.match(text)) and not text.endswith("tölur.")


def drop_ragged_zeros(values: list[list], fig_id: str, names: list[str]) -> None:
    """Turn zeros that stand in for not-yet-published values into gaps (in place).

    When one country has a newer observation than the others, the export fills the others'
    last cells with 0 (X-9: 106, 106, 0). A short run of zeros (at most 3) at the end of a
    series counts as missing when another series has data on those dates and the series'
    last value is far from zero (at least a quarter of its largest value). Zeros inside a
    series are left alone: months without issuance or lending are real zeros.
    """
    for si, s in enumerate(values):
        end = len(s)
        while end and s[end - 1] is None:
            end -= 1
        start = end
        while start and s[start - 1] == 0:
            start -= 1
        if start == end or end - start > 3 or start == 0:
            continue
        prev = next((v for v in reversed(s[:start]) if v is not None), None)
        top = max(abs(v) for v in s if v is not None)
        others = all(any(o[i] not in (None, 0) for oi, o in enumerate(values) if oi != si) for i in range(start, end))
        if prev is not None and abs(prev) >= 0.25 * top and others:
            for i in range(start, end):
                s[i] = None
            warn(fig_id, f"{names[si]}: last {end - start} zero(s) treated as missing")


def read_meta(rows: list[tuple]) -> tuple[dict, int]:
    """Metadata from the labelled rows (Fs., Ufs., Nm., H., Vá., Há., Ath.) or, without labels, by position."""
    meta = {k: None for k in META_ROWS.values()}
    meta["latest"] = None
    last = 9
    found = False
    for i, row in enumerate(rows[:16]):
        key = META_LABELS.get(re.sub(r"[\s.]", "", str(row[0] or "")).lower()) if row else None
        if key:
            found = True
            meta[key] = clean_text(row[1]) if len(row) > 1 else None
            last = max(last, i)
    if not found:
        # No labels (chapters VII–IX): place the rows relative to the "Heimild" line.
        cell = lambda i: clean_text(rows[i][1]) if i < len(rows) and len(rows[i]) > 1 else None
        src = next((i for i in range(3, 12) if (cell(i) or "").lower().startswith("heimild")), 6)
        rest = [cell(i) for i in range(3, src + 3) if i != src and cell(i)]
        meta["title"] = rest.pop(0) if rest else None
        units = [t for t in rest if looks_like_unit(t)]
        notes = [t for t in rest if not looks_like_unit(t) and not t.lower().startswith("nýjast")]
        meta["latest"] = next((t for t in rest if t.lower().startswith("nýjast")), None)
        meta["sources"] = cell(src)
        meta["unit"] = units[0] if units else None
        meta["unitRight"] = units[1] if len(units) > 1 else None
        meta["note"] = " ".join(notes) or None
        last = src + (2 if cell(src + 2) else 1 if cell(src + 1) else 0)
    return meta, last


def parse_sheet(rows: list[tuple], fig_id: str, chapter: str, toc_title: str | None, override: dict) -> tuple[dict, dict] | None:
    meta, last_meta = read_meta(rows)
    if meta.get("note"):
        m = re.search(r"\s*(Nýjast:[^.]*\d{4}\.?)\s*$", meta["note"])
        if m:
            meta["latest"] = meta.get("latest") or m.group(1).rstrip(".")
            meta["note"] = meta["note"][: m.start()].strip() or None
    for key in ("unitRight", "unit", "subtitle"):
        if meta[key] and meta[key].lower().startswith("nýjast"):
            meta["latest"] = meta[key]
            meta[key] = None
    title = re.sub(r"(?<=[a-záéíóúýþæöð])\d$", "", meta["title"] or "") or toc_title
    if not title:
        warn(fig_id, "no title")
        title = fig_id

    start = override.get("dataRow")
    if start is None:
        start = find_data_start(rows, last_meta + 1)
    else:
        start -= 1
    if start is None:
        # Some sheets are published without data (licensed sources); keep them listed.
        warn(fig_id, "no data in sheet, listed as unavailable")
        header = next((list(r) for r in rows[last_meta + 1 :] if r and not blank(r) and any(is_string(v) for v in r[1:])), [])
        names = [split_axis(label(v)) for v in header[1:] if clean_text(v)]
        entry = base_entry(fig_id, chapter, title, meta)
        entry.update(xType="none", frequency=None, coverage=None, points=0, series=[{"name": n, "axis": a} for n, a in names])
        entry["unavailable"] = meta["comment"] or "Gögn eru ekki birt."
        return entry, None
    h = start - 1
    while h > last_meta and blank(rows[h]):
        h -= 1
    header = list(rows[h]) if h > last_meta else []
    group = list(rows[h - 1]) if h - 1 > last_meta and not blank(rows[h - 1]) else []

    data = [list(r) for r in rows[start:] if not blank(r)]
    # Daily market series repeat a date now and then; keep the last observation.
    if all(isinstance(r[0], dt.datetime) for r in data):
        last = {r[0]: i for i, r in enumerate(data)}
        if len(last) != len(data):
            warn(fig_id, f"{len(data) - len(last)} duplicate dates dropped")
            data = [r for i, r in enumerate(data) if last[r[0]] == i]
    if sum(isinstance(v, dt.datetime) for v in header[1:]) >= 3 and all(is_string(r[0]) for r in data):
        header, data = transpose_block(header, data)
        group = []
    width = max(len(r) for r in data)
    header += [None] * (width - len(header))
    group += [None] * (width - len(group))
    for r in data:
        r += [None] * (width - len(r))

    # Column B holds sparse axis labels ("'98") when it has no header and no numbers.
    b_nums = [to_number(r[1]) for r in data if width > 1 and to_number(r[1]) is not None]
    yearish = all(v.is_integer() and (v < 100 or 1900 <= v <= 2100) for v in b_nums)
    b_is_label = width > 1 and (
        (clean_text(header[1]) is None and (len(b_nums) < 0.3 * len(data) or yearish))
        or (bool(b_nums) and yearish and len(b_nums) < 0.3 * len(data))
    )
    first_series_col = 2 if b_is_label else 1
    cols, unnamed = [], []
    for c in range(first_series_col, width):
        values = [to_number(r[c]) for r in data]
        name = label(header[c]) if header[c] is not None else ""
        if not any(v is not None for v in values):
            continue
        if not name:
            unnamed.append((c, values))
            continue
        cols.append((c, name, values))
    if unnamed and not cols and len(unnamed) == 1:
        cols = [(unnamed[0][0], title, unnamed[0][1])]
    elif unnamed:
        warn(fig_id, f"columns {', '.join(str(c + 1) for c, _ in unnamed)} have values but no header, skipped")
    if not cols:
        warn(fig_id, "no numeric series, skipped")
        return None

    # Names: strip axis marker, use group row when names clash.
    names = [split_axis(n) for _, n, _ in cols]
    plain = [n for n, _ in names]
    if len(set(plain)) != len(plain) and any(clean_text(g) for g in group):
        full_row = all(clean_text(group[c]) for c, _, _ in cols)
        filled, cur = [], None
        for c in range(width):
            cur = clean_text(group[c]) or cur
            filled.append(cur)
        if not filled[cols[0][0]]:
            # labels centred over their columns (III-19): take the nearest label
            marks = [c for c in range(width) if clean_text(group[c])]
            filled = [clean_text(group[min(marks, key=lambda m: (abs(m - c), m))]) if marks else None for c in range(width)]
        renamed = []
        for (c, _, _), (n, axis) in zip(cols, names):
            if full_row:
                g, g_axis = split_axis(clean_text(group[c]))
                renamed.append((g, "right" if "right" in (axis, g_axis) else "left"))
            else:
                renamed.append((f"{filled[c]}: {n}" if filled[c] else n, axis))
        names = renamed
        plain = [n for n, _ in names]
    if len(set(plain)) != len(plain):
        # same name on both axes (II-2): keep the axis in the name, as the report's legends do
        counts = Counter(plain)
        names = [(f"{n} ({'h' if a == 'right' else 'v'}. ás)" if counts[n] > 1 and len({x[1] for x in names if x[0] == n}) > 1 else n, a) for n, a in names]
        plain = [n for n, _ in names]
    if len(set(plain)) != len(plain):
        seen = Counter()
        fixed = []
        for n, axis in names:
            seen[n] += 1
            fixed.append((n if seen[n] == 1 else f"{n} ({seen[n]})", axis))
        names = fixed
        warn(fig_id, "duplicate series names numbered")

    if override.get("latestBlock"):
        # long format (VII-29): one block of rows per date; show the latest block by column B
        latest = max(r[0] for r in data)
        keep_rows = [i for i, r in enumerate(data) if r[0] == latest]
        data = [data[i] for i in keep_rows]
        cols = [(c, n, [v[i] for i in keep_rows]) for c, n, v in cols]
        for r in data:
            r[0] = r[1]
        meta["subtitle"] = meta["subtitle"] or f"Staða {label(latest)}"
    if override.get("dropColumns"):
        keep = [i for i, (n, _) in enumerate(names) if n not in override["dropColumns"]]
        cols, names = [cols[i] for i in keep], [names[i] for i in keep]
    a_col = [r[0] for r in data]
    b_col = [r[1] for r in data] if b_is_label else None
    if override.get("x") == "category":
        x_type, x, freq = "category", [label(v) for v in a_col], None
    else:
        x_type, x, freq = resolve_x(a_col, b_col)
    if x_type == "category":
        warn(fig_id, f"category axis ({x[0]!r} … {x[-1]!r}, {len(x)} labels)")
        if any(not v for v in x):
            # rows without a label in column A: borrow column B's display label if any
            x = [v or (label(b) if b_col else "") for v, b in zip(x, b_col or [None] * len(x))]

    values = [[round_value(v) for v in vals] for _, _, vals in cols]
    if override.get("transpose") and x_type == "category":
        # series are years and rows are items (IX-1): draw the years along the x-axis
        names, x = [(c, "left") for c in x], [n for n, _ in names]
        values = [list(col) for col in zip(*values)]
    series_meta = [{"name": n, "axis": axis.split("-")[0], **({"inverse": True} if axis.endswith("inverse") else {})} for n, axis in names]

    if x_type == "time" and not override.get("keepZeros"):
        drop_ragged_zeros(values, fig_id, [n for n, _ in names])

    # trim trailing rows where every series is empty (forecast placeholders etc.)
    keep = len(x)
    while keep > 1 and all(v[keep - 1] is None for v in values):
        keep -= 1
    x = x[:keep]
    values = [v[:keep] for v in values]

    entry = base_entry(fig_id, chapter, title, meta)
    entry.update({
        "xType": x_type,
        "frequency": freq,
        "coverage": [x[0], x[-1]] if x_type == "time" else None,
        "points": len(x),
        "series": series_meta,
    })
    if x_type == "category":
        entry["categories"] = x
    if any(s["axis"] == "right" for s in series_meta) and not entry["unitRight"]:
        entry["unitRight"] = entry["unit"]
    payload = {"id": fig_id, "x": x, "series": values}
    return entry, payload


def base_entry(fig_id: str, chapter: str, title: str, meta: dict) -> dict:
    return {
        "id": fig_id,
        "chapterId": chapter,
        "title": title,
        "subtitle": meta["subtitle"],
        "note": meta["note"],
        "sources": parse_sources(meta["sources"]),
        "unit": meta["unit"] or "",
        "unitRight": meta["unitRight"],
        "latest": meta.get("latest"),
    }


def figure_id(sheet_name: str, cell_label) -> str:
    s = clean_text(cell_label) or ""
    m = re.search(r"\b([IVX]+-\d+)\b", s) or re.search(r"\b([IVX]+-\d+)\b", sheet_name)
    return m.group(1) if m else sheet_name.rstrip(". ")


def read_supplement(path: Path) -> list[list]:
    """CSV with a header row, then one row per observation: date (YYYY-MM-DD) or label, then one value per series."""
    import csv

    text = path.read_text("utf-8-sig")
    delim = ";" if text.count(";") > text.count(",") else ","
    out = []
    for i, row in enumerate(csv.reader(text.splitlines(), delimiter=delim)):
        if i == 0 or not row or not row[0].strip():
            continue
        x = row[0].strip()
        try:
            x = dt.datetime.strptime(x[:10], "%Y-%m-%d")
        except ValueError:
            pass
        out.append([x, *[to_number(v.replace(",", ".") if delim == ";" else v) for v in row[1:]]])
    return out


def merge_supplement(rows: list[tuple], data: list[list]) -> list[tuple]:
    """Put supplement values into an empty sheet under its own header row, so the normal parser reads it."""
    meta_end = read_meta(rows)[1]
    h = next((i for i in range(meta_end + 1, len(rows)) if rows[i] and any(is_string(v) for v in rows[i][1:])), None)
    if h is None:
        return rows
    header = list(rows[h])
    pad = [None] if len(header) > 1 and clean_text(header[1]) is None else []
    return [*rows[: h + 1], *[tuple([r[0], *pad, *r[1:]]) for r in data]]


def parse_workbook(path: Path, chapter: str, overrides: dict, supplements: Path | None = None) -> tuple[str | None, list, dict]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    toc: dict[str, str] = {}
    chapter_title = None
    if "Myndayfirlit" in wb.sheetnames:
        for i, row in enumerate(wb["Myndayfirlit"].iter_rows(values_only=True)):
            if i == 1 and row and row[0]:
                chapter_title = re.sub(rf"^{chapter}\s+", "", clean_text(row[0]) or "")
            if row and row[0] and len(row) > 1 and row[1] and re.fullmatch(r"[IVX]+-\d+\.?", str(row[0]).strip()):
                toc[str(row[0]).strip().rstrip(".")] = re.sub(r"(\D)\d$", r"\1", clean_text(row[1]) or "")
    entries, payloads = [], {}
    for name in wb.sheetnames:
        if name == "Myndayfirlit":
            continue
        rows = list(wb[name].iter_rows(values_only=True))
        fid = figure_id(name, rows[2][1] if len(rows) > 2 and len(rows[2]) > 1 else None)
        if overrides.get(fid, {}).get("skip"):
            warn(fid, "skipped by override")
            continue
        supplement = supplements / f"{fid}.csv" if supplements else None
        if supplement and supplement.exists():
            rows = merge_supplement(rows, read_supplement(supplement))
            warn(fid, f"values from {supplement.name}")
        try:
            res = parse_sheet(rows, fid, chapter, toc.get(fid), overrides.get(fid, {}))
        except Exception as exc:  # keep going; report at the end
            warn(fid, f"failed to parse: {exc!r}")
            continue
        if res:
            entries.append(res[0])
            if res[1]:
                payloads[fid] = res[1]
    wb.close()
    return chapter_title, entries, payloads


# --------------------------------------------------------------------------- #
#  Main
# --------------------------------------------------------------------------- #


def write_json(path: Path, obj) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(obj, ensure_ascii=False, separators=(",", ":")), "utf-8")


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--sources", type=Path, default=HERE / "sources.json")
    ap.add_argument("--overrides", type=Path, default=HERE / "overrides.json")
    ap.add_argument("--out", type=Path, default=ROOT / "public" / "data")
    ap.add_argument("--supplements", type=Path, default=HERE / "supplements", help="CSV files (<figure id>.csv) for sheets published without data")
    ap.add_argument("--local", type=Path, help="directory with already downloaded workbooks (file name = URL basename)")
    ap.add_argument("--check", action="store_true", help="only check the server for changes (exit 10 if changed)")
    args = ap.parse_args()

    sources = json.loads(args.sources.read_text("utf-8"))
    if args.check:
        return check(sources, args.out / "manifest.json")
    overrides = json.loads(args.overrides.read_text("utf-8")) if args.overrides.exists() else {}
    overrides = {k: v for k, v in overrides.items() if not k.startswith("_")}

    tmp = None
    if args.local:
        src_dir = args.local
    else:
        tmp = tempfile.TemporaryDirectory(prefix="hagvisar-")
        src_dir = Path(tmp.name)

    manifest = {"release": sources["release"], "releaseLabel": sources.get("releaseLabel"), "generatedAt": dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds"), "workbooks": {}}
    chapters, figures = [], []
    fig_dir = args.out / "figures"
    fig_dir.mkdir(parents=True, exist_ok=True)
    written = set()

    jobs = {ch: sources["workbooks"][ch] for ch in CHAPTERS if sources["workbooks"].get(ch)}
    paths = {ch: src_dir / url.rsplit("/", 1)[1] for ch, url in jobs.items()}
    if args.local:
        infos = {}
        for ch in jobs:
            if not paths[ch].exists():
                paths[ch] = next(iter(sorted(src_dir.glob(f"{ch}-*.xlsx"))), paths[ch])
            body = paths[ch].read_bytes()
            infos[ch] = {"etag": None, "lastModified": None, "sha256": hashlib.sha256(body).hexdigest(), "bytes": len(body)}
    else:
        with ThreadPoolExecutor(max_workers=5) as pool:
            infos = dict(zip(jobs, pool.map(lambda ch: download(jobs[ch], paths[ch]), jobs)))

    for ch, url in jobs.items():
        path, info = paths[ch], infos[ch]
        title, entries, payloads = parse_workbook(path, ch, overrides, args.supplements)
        updated = None
        if info.get("lastModified"):
            updated = dt.datetime.strptime(info["lastModified"], "%a, %d %b %Y %H:%M:%S %Z").date().isoformat()
        for e in entries:
            e["updatedAt"] = updated or dt.date.today().isoformat()
        for fid, p in payloads.items():
            write_json(fig_dir / f"{fid}.json", p)
            written.add(f"{fid}.json")
        chapters.append({"id": ch, "title": title})
        figures.extend(entries)
        manifest["workbooks"][ch] = {"url": url, **info, "figures": [e["id"] for e in entries]}
        print(f"{ch:>4}: {len(entries):3d} figures  ({title})")

    for stale in fig_dir.glob("*.json"):
        if stale.name not in written:
            stale.unlink()

    write_json(args.out / "catalogue.json", {"release": sources["release"], "releaseLabel": sources.get("releaseLabel"), "generatedAt": manifest["generatedAt"], "chapters": chapters, "figures": figures})
    (args.out / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), "utf-8")
    if tmp:
        tmp.cleanup()

    print(f"\n{len(figures)} figures written to {args.out}")
    if warnings:
        print(f"\n{len(warnings)} notes:")
        for w in warnings:
            print("  " + w)
    return 0


if __name__ == "__main__":
    sys.exit(main())
